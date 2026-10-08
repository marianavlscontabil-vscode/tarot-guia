# -*- coding: utf-8 -*-
"""
Gera combinacoes_mistas_5_cartas_tarot_egipcio.xlsx
Tarô Legítimo Egípcio — TODAS as C(40,5) = 658.008 combinações mistas de 5 cartas
(22 Arcanos Maiores + 16 Menores + 2 Especiais), sem repetição e sem ordem.

Requisitos:  pip install openpyxl
Execução:    python gerar_planilha_tarot.py
Tempo estimado: 1 a 3 minutos. Ao final, imprime o bloco
"RESUMO PARA CONFERÊNCIA" com as contagens para validar.
"""

from itertools import combinations
from collections import Counter
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.cell import WriteOnlyCell

ARQUIVO = "combinacoes_mistas_5_cartas_tarot_egipcio.xlsx"

# ============================================================
# 1. BARALHO (40 cartas)
# (nome, grupo, naipe, palavra_chave, essencia, polaridade, significado)
# ============================================================
BRUTO = [
    # --- 22 Arcanos Maiores ---
    ("0. O Louco", "Maior", None, "Novo começo", "um novo começo cheio de possibilidades", "Positivo", "Início de ciclo, liberdade, risco calculado, fé no desconhecido"),
    ("I. O Mago", "Maior", None, "Iniciativa", "iniciativa e recursos para agir", "Positivo", "Habilidade, vontade, recursos para agir, comunicação"),
    ("II. A Sacerdotisa", "Maior", None, "Intuição", "intuição e saber interior", "Neutro", "Saber oculto, mistério, escutar o interior"),
    ("III. A Imperatriz", "Maior", None, "Abundância", "abundância, cuidado e criação", "Positivo", "Fertilidade, cuidado, criatividade, nutrição"),
    ("IV. O Imperador", "Maior", None, "Estrutura", "estrutura, ordem e autoridade", "Positivo", "Autoridade, ordem, estabilidade, disciplina"),
    ("V. O Hierofante", "Maior", None, "Tradição", "tradição, ensino e orientação", "Neutro", "Ensino, normas, espiritualidade, conselho"),
    ("VI. Os Amantes", "Maior", None, "Escolha", "escolhas do coração e união", "Positivo", "Amor, decisão de coração, união, valores"),
    ("VII. O Carro", "Maior", None, "Avanço", "determinação e avanço vitorioso", "Positivo", "Determinação, vitória, controle da direção"),
    ("VIII. A Força", "Maior", None, "Coragem", "coragem e domínio interior", "Positivo", "Domínio interior, paciência, brandura firme"),
    ("IX. O Eremita", "Maior", None, "Recolhimento", "recolhimento e busca interior", "Neutro", "Reflexão, busca interior, sabedoria solitária"),
    ("X. A Roda da Fortuna", "Maior", None, "Ciclos", "viradas de sorte e ciclos do destino", "Neutro", "Sorte, viradas, pontos de virada do destino"),
    ("XI. A Justiça", "Maior", None, "Equilíbrio", "verdade, justiça e consequências", "Neutro", "Verdade, consequências, decisões justas, contratos"),
    ("XII. O Enforcado", "Maior", None, "Pausa", "pausa, sacrifício e nova perspectiva", "Neutro", "Suspensão, sacrifício, ver de outro ângulo"),
    ("XIII. A Morte", "Maior", None, "Transformação", "transformação e fim de ciclo", "Desafiador", "Fim de ciclo, renovação inevitável, transição"),
    ("XIV. A Temperança", "Maior", None, "Moderação", "moderação, equilíbrio e cura", "Positivo", "Equilíbrio, paciência, harmonia, cura"),
    ("XV. O Diabo", "Maior", None, "Apego", "apego, tentação e vínculos intensos", "Desafiador", "Dependência, tentação, materialismo, paixões cegas"),
    ("XVI. A Torre", "Maior", None, "Ruptura", "ruptura súbita e queda de estruturas", "Desafiador", "Colapso súbito, revelação, queda de estruturas falsas"),
    ("XVII. A Estrela", "Maior", None, "Esperança", "esperança, cura e inspiração", "Positivo", "Cura, fé, inspiração, serenidade após a tempestade"),
    ("XVIII. A Lua", "Maior", None, "Ilusão", "ilusões, medos e caminhos obscuros", "Desafiador", "Medos, incertezas, sonhos, caminhos obscuros"),
    ("XIX. O Sol", "Maior", None, "Êxito", "sucesso, vitalidade e clareza", "Positivo", "Vitalidade, sucesso, clareza, alegria"),
    ("XX. O Julgamento", "Maior", None, "Despertar", "despertar, balanço e renascimento", "Positivo", "Balanço de vida, chamado, renascimento, perdão"),
    ("XXI. O Mundo", "Maior", None, "Conclusão", "conclusão plena e realização", "Positivo", "Fechamento pleno, realização, viagem, integridade"),
    # --- 16 Arcanos Menores (Paus, Copas, Ouros, Espadas × Ás, Valete, Rainha, Rei) ---
    ("Ás de Paus", "Menor", "paus", "Iniciativa", "uma nova energia de ação e começo de projetos", "Positivo", "Nova iniciativa, faísca de projeto, impulso criativo no trabalho"),
    ("Valete de Paus", "Menor", "paus", "Entusiasmo", "entusiasmo em formação e boas notícias no trabalho", "Neutro", "Energia jovem, aprendizado prático, recado sobre empreendimentos"),
    ("Rainha de Paus", "Menor", "paus", "Determinação", "determinação calorosa e liderança acolhedora", "Positivo", "Liderança carismática, coragem, energia que inspira os outros"),
    ("Rei de Paus", "Menor", "paus", "Liderança", "liderança, visão e autoridade empreendedora", "Positivo", "Visão empreendedora, comando firme, sucesso por iniciativa"),
    ("Ás de Copas", "Menor", "copas", "Novo amor", "um novo sentimento que nasce", "Positivo", "Amor que nasce, abertura emocional, intuição aflorada"),
    ("Valete de Copas", "Menor", "copas", "Recado do coração", "um recado romântico ou sensibilidade jovem", "Positivo", "Mensagem afetiva, romantismo, sensibilidade e boas novas emocionais"),
    ("Rainha de Copas", "Menor", "copas", "Empatia", "empatia, cuidado e intuição emocional", "Positivo", "Compreensão, acolhimento, sabedoria do coração"),
    ("Rei de Copas", "Menor", "copas", "Equilíbrio emocional", "equilíbrio emocional e generosidade", "Positivo", "Maturidade afetiva, calma, generosidade e diplomacia"),
    ("Ás de Ouros", "Menor", "ouros", "Nova oportunidade", "uma nova oportunidade material", "Positivo", "Oportunidade de ganho, semente de prosperidade, oferta concreta"),
    ("Valete de Ouros", "Menor", "ouros", "Aplicação", "aplicação diligente e aprendizado prático", "Neutro", "Estudo, dedicação, primeiro passo na construção material"),
    ("Rainha de Ouros", "Menor", "ouros", "Praticidade", "praticidade generosa e segurança material", "Positivo", "Gestão sensata, prosperidade compartilhada, cuidado com o lar"),
    ("Rei de Ouros", "Menor", "ouros", "Sucesso material", "sucesso material e estabilidade", "Positivo", "Conquista financeira, solidez, mestria nos negócios"),
    ("Ás de Espadas", "Menor", "espadas", "Verdade", "uma verdade que corta e clareza brutal", "Neutro", "Clareza, decisão definitiva, verdade que pode ferir mas liberta"),
    ("Valete de Espadas", "Menor", "espadas", "Vigilância", "vigilância e palavras afiadas", "Desafiador", "Conflito iminente, fofoca, necessidade de cautela com palavras"),
    ("Rainha de Espadas", "Menor", "espadas", "Lucidez", "lucidez fria e independência", "Neutro", "Percepção aguda, distanciamento emocional, julgamento justo porém frio"),
    ("Rei de Espadas", "Menor", "espadas", "Julgamento", "autoridade intelectual e julgamento severo", "Neutro", "Autoridade racional, decisões duras mas justas, ética"),
    # --- 2 Cartas Especiais ---
    ("A Riqueza", "Especial", None, "Prosperidade", "ganho material e sorte", "Positivo", "Ganho material, abundância iminente, sorte nos negócios"),
    ("A Vida", "Especial", None, "Vitalidade", "vitalidade, saúde e renovação", "Positivo", "Saúde, energia vital, renovação e recomeço"),
]

assert len(BRUTO) == 40, "O baralho deve ter exatamente 40 cartas"

nomes   = [c[0] for c in BRUTO]
grupos  = [c[1] for c in BRUTO]
naipes  = [c[2] for c in BRUTO]
pals    = [c[3] for c in BRUTO]
ess     = [c[4] for c in BRUTO]
pols    = [c[5] for c in BRUTO]

p_pos = [1 if p == "Positivo" else 0 for p in pols]
p_des = [1 if p == "Desafiador" else 0 for p in pols]

# Pontuação por área (1 ponto por carta): chaves da área OU naipe correspondente
CHAVES_AMOR  = {"VI. Os Amantes", "III. A Imperatriz", "XVII. A Estrela", "XV. O Diabo", "XVIII. A Lua"}
CHAVES_TRAB  = {"I. O Mago", "IV. O Imperador", "VII. O Carro", "VIII. A Força", "XI. A Justiça"}
CHAVES_DIN   = {"III. A Imperatriz", "X. A Roda da Fortuna", "XIX. O Sol", "XV. O Diabo", "XVI. A Torre", "A Riqueza"}

pt_amor = [1 if (nomes[i] in CHAVES_AMOR or naipes[i] == "copas") else 0 for i in range(40)]
pt_trab = [1 if (nomes[i] in CHAVES_TRAB or naipes[i] == "paus")  else 0 for i in range(40)]
pt_din  = [1 if (nomes[i] in CHAVES_DIN  or naipes[i] == "ouros") else 0 for i in range(40)]

DESTAQUES = {
    "0. O Louco": "Destaque: O Louco marca um recomeço inesperado.",
    "XIII. A Morte": "Destaque: A Morte marca transformação inevitável.",
    "XVI. A Torre": "Destaque: A Torre marca ruptura libertadora.",
    "X. A Roda da Fortuna": "Destaque: A Roda da Fortuna marca virada de fase.",
    "XIX. O Sol": "Destaque: O Sol marca sucesso visível.",
    "XXI. O Mundo": "Destaque: O Mundo marca fechamento de ciclo.",
    "A Riqueza": "Destaque: A Riqueza marca ganho material iminente.",
    "A Vida": "Destaque: A Vida marca renovação e saúde.",
}
dest_txt = [DESTAQUES.get(n, "") for n in nomes]

FRASES = {
    "Muito favorável": "As cinco energias convergem para um resultado próspero e maduro.",
    "Favorável": "O conjunto favorece o avanço, desde que se mantenha consciência nos detalhes.",
    "Misto": "Há avanço com pontos de atenção: a energia desafiadora pede cautela e honestidade.",
    "Neutro": "O momento pede observação e equilíbrio antes de grandes decisões.",
    "Tenso": "O conjunto aponta tensões a enfrentar; as energias positivas indicam o caminho de saída.",
    "Desafiador": "O conjunto aponta um período de provas: encarar os obstáculos com lucidez é o caminho.",
}

CONSELHO = {
    ("Amor", "forte"):    "O amor está em destaque nesta tiragem: cultive os vínculos, expresse sentimentos e cuide das parcerias.",
    ("Amor", "presente"): "O coração tem papel na questão: escute os sentimentos antes de tomar decisões.",
    ("Amor", "neutro"):   "O foco não está no amor agora: mantenha os vínculos em equilíbrio, sem grandes mudanças.",
    ("Trabalho", "forte"):    "A carreira domina a cena: organize planos, aja com disciplina e busque reconhecimento.",
    ("Trabalho", "presente"): "O trabalho aparece na questão: avance com método e clareza nos objetivos.",
    ("Trabalho", "neutro"):   "O momento pede menos foco profissional: mantenha a rotina e observe as oportunidades.",
    ("Dinheiro", "forte"):    "As finanças estão em jogo: planeje gastos, aproveite oportunidades e evite excessos.",
    ("Dinheiro", "presente"): "O dinheiro entra na leitura: avalie investimentos e gastos com atenção redobrada.",
    ("Dinheiro", "neutro"):   "As finanças estão em equilíbrio: mantenha o controle habitual e evite riscos desnecessários.",
}

def nivel(pts):
    return "forte" if pts >= 2 else ("presente" if pts == 1 else "neutro")

# ============================================================
# 2. WORKBOOK (modo write_only — uma única passada de escrita)
# ============================================================
wb = Workbook(write_only=True)
wb._named_styles[0].font = Font(name="Arial", size=11)  # fonte padrão do arquivo

def cab(ws, titulos):
    celulas = []
    for t in titulos:
        c = WriteOnlyCell(ws, value=t)
        c.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celulas.append(c)
    ws.append(celulas)

# ---------- ABA 1: Guia das Cartas ----------
ws1 = wb.create_sheet("Guia das Cartas")
cab(ws1, ["Grupo", "Carta", "Palavra-chave", "Essência", "Polaridade", "Significado"])
for c in BRUTO:
    ws1.append([c[1], c[0], c[3], c[4], c[5], c[6]])
ws1.append([])
ws1.append(["Nota: Polaridades dos Menores seguem leitura convencional: ases e figuras de Paus, Copas e Ouros tendem ao positivo; Espadas carrega a tensão do naipe. As Especiais funcionam como trunfos: quando saem, dominam a leitura."])
for col, w in zip("ABCDEF", (10, 22, 20, 45, 12, 60)):
    ws1.column_dimensions[col].width = w
ws1.freeze_panes = "A2"
ws1.auto_filter.ref = "A1:F41"

# ---------- ABA 2: Posições da Cruz de 5 ----------
ws2 = wb.create_sheet("Posições da Cruz de 5")
cab(ws2, ["Posição", "Nome", "Se cair um Arcano Maior", "Se cair um Arcano Menor", "Se cair uma Especial"])
ws2.append([1, "Situação", "A questão é regida por uma força do destino; o tema é grande e inevitável", "A situação é prática e do cotidiano; o naipe indica a área (Paus trabalho, Copas amor, Ouros dinheiro, Espadas conflitos)", "A questão gira em torno de ganho (Riqueza) ou renovação (Vida)"])
ws2.append([2, "Desafio", "O obstáculo é uma lição de vida; não se resolve por força bruta", "Obstáculo concreto do dia a dia; figuras da corte apontam a pessoa que dificulta", "O desafio é não desperdiçar a oportunidade ou a energia vital"])
ws2.append([3, "Base (inconsciente)", "Raiz profunda: um padrão de destino ou crença maior move a questão", "Raiz prática: hábitos, rotinas ou pessoas próximas sustentam a situação", "No fundo, a questão fala de segurança material ou vitalidade"])
ws2.append([4, "Passado recente", "Um evento marcante e inevitável moldou o momento", "Acontecimentos cotidianos recentes explicam o presente", "Ganho ou recuperação recente influencia o agora"])
ws2.append([5, "Resultado", "Desfecho nas mãos do destino; peso máximo na leitura", "Desfecho prático, moldável pelas suas ações", "Desfecho próspero (Riqueza) ou de recomeço (Vida)"])
ws2.append([])
cab(ws2, ["Cartas de peso (Maiores+Especiais)", "Menores", "Combinações", "Leitura da proporção"])
ws2.append([0, 5, 4368,   'Tiragem totalmente prática: tudo está nas suas mãos; nada está "escrito"'])
ws2.append([1, 4, 43680,  "Uma força maior pontua o cotidiano: um sinal dentro da rotina"])
ws2.append([2, 3, 154560, "Equilíbrio entre destino e ação própria: o cenário mais comum e legível"])
ws2.append([3, 2, 242880, "O destino predomina: momento decisivo, o cotidiano acompanha"])
ws2.append([4, 1, 170016, "Quase tudo nas mãos de forças maiores: adapte-se em vez de forçar"])
ws2.append([5, 0, 42504,  "Tiragem de destino puro: fase de virada de vida"])
ws2.append([])
ws2.append(["Total: 658.008 combinações — C(40,5). Esta planilha lista as combinações sem ordem; para ler a Cruz de 5 com posições, aplique o significado de cada posição à carta que cair nela."])
for col, w in zip("ABCDE", (30, 22, 60, 60, 55)):
    ws2.column_dimensions[col].width = w

# ---------- ABA 3: Combinações Mistas (5 cartas) ----------
ws3 = wb.create_sheet("Combinações Mistas (5 cartas)")
cab(ws3, ["Nº", "Carta 1", "Carta 2", "Carta 3", "Carta 4", "Carta 5",
          "Grupos", "Palavras-chave", "Tom geral", "Leitura da combinação",
          "Conselho — Amor", "Conselho — Trabalho", "Conselho — Dinheiro"])

cont_tom, cont_amor, cont_trab, cont_din = Counter(), Counter(), Counter(), Counter()
append3 = ws3.append
n = 0

for combo in combinations(range(40), 5):
    n += 1
    P = p_pos[combo[0]] + p_pos[combo[1]] + p_pos[combo[2]] + p_pos[combo[3]] + p_pos[combo[4]]
    D = p_des[combo[0]] + p_des[combo[1]] + p_des[combo[2]] + p_des[combo[3]] + p_des[combo[4]]
    if D == 0:
        tom = "Muito favorável" if P >= 3 else "Favorável"
    elif D == 1:
        tom = "Misto"
    elif D == 2:
        tom = "Tenso" if P >= 1 else "Neutro"
    else:
        tom = "Desafiador"
    cont_tom[tom] += 1

    a = pt_amor[combo[0]] + pt_amor[combo[1]] + pt_amor[combo[2]] + pt_amor[combo[3]] + pt_amor[combo[4]]
    t = pt_trab[combo[0]] + pt_trab[combo[1]] + pt_trab[combo[2]] + pt_trab[combo[3]] + pt_trab[combo[4]]
    d = pt_din[combo[0]] + pt_din[combo[1]] + pt_din[combo[2]] + pt_din[combo[3]] + pt_din[combo[4]]
    na, nt, nd = nivel(a), nivel(t), nivel(d)
    cont_amor[na] += 1; cont_trab[nt] += 1; cont_din[nd] += 1

    dest = " ".join(x for x in (dest_txt[i] for i in combo) if x)
    leitura = (f"Energias de {ess[combo[0]]}, combinadas com {ess[combo[1]]}, {ess[combo[2]]}, "
               f"{ess[combo[3]]} e {ess[combo[4]]}. {FRASES[tom]}" + (f" {dest}" if dest else ""))

    append3([n, nomes[combo[0]], nomes[combo[1]], nomes[combo[2]], nomes[combo[3]], nomes[combo[4]],
             " · ".join(grupos[i] for i in combo),
             " · ".join(pals[i] for i in combo),
             tom, leitura,
             CONSELHO[("Amor", na)], CONSELHO[("Trabalho", nt)], CONSELHO[("Dinheiro", nd)]])

ws3.freeze_panes = "A2"
ws3.auto_filter.ref = f"A1:M{n + 1}"
for col, w in zip("ABCDEFGHIJKLM", (8, 20, 20, 20, 20, 20, 28, 50, 15, 90, 55, 55, 55)):
    ws3.column_dimensions[col].width = w

wb.save(ARQUIVO)
print(f"Arquivo salvo: {ARQUIVO} — {n} combinações escritas.")

# ============================================================
# 3. VERIFICAÇÃO LEVE (sem recarregar as 658 mil linhas)
# ============================================================
wb2 = load_workbook(ARQUIVO, read_only=True)
ws2v = wb2["Combinações Mistas (5 cartas)"]
print(f"Linhas da aba de combinações (com cabeçalho): {ws2v.max_row} (esperado 658009)")
primeiras = list(ws2v.iter_rows(max_row=2, values_only=True))
print("Amostra linha 1:", primeiras[1])
wb2.close()

ESPERADO_TOM = {"Muito favorável": 270710, "Favorável": 53922, "Misto": 261800,
                "Neutro": 1650, "Tenso": 63800, "Desafiador": 6126}
ESPERADO_AMOR = {"forte": 204912, "presente": 283185, "neutro": 169911}
ESPERADO_TRAB = {"forte": 204912, "presente": 283185, "neutro": 169911}
ESPERADO_DIN  = {"forte": 241452, "presente": 274050, "neutro": 142506}

print("\n===== RESUMO PARA CONFERÊNCIA =====")
print(f"Total de combinações: {n} (esperado 658008) -> {'OK' if n == 658008 else 'DIVERGÊNCIA'}")
print("--- Tons ---")
for k, v in ESPERADO_TOM.items():
    print(f"{k}: obtido {cont_tom[k]} | esperado {v} -> {'OK' if cont_tom[k] == v else 'DIVERGÊNCIA'}")
print("--- Conselho Amor ---")
for k, v in ESPERADO_AMOR.items():
    print(f"{k}: obtido {cont_amor[k]} | esperado {v} -> {'OK' if cont_amor[k] == v else 'DIVERGÊNCIA'}")
print("--- Conselho Trabalho ---")
for k, v in ESPERADO_TRAB.items():
    print(f"{k}: obtido {cont_trab[k]} | esperado {v} -> {'OK' if cont_trab[k] == v else 'DIVERGÊNCIA'}")
print("--- Conselho Dinheiro ---")
for k, v in ESPERADO_DIN.items():
    print(f"{k}: obtido {cont_din[k]} | esperado {v} -> {'OK' if cont_din[k] == v else 'DIVERGÊNCIA'}")
print("===== FIM DA CONFERÊNCIA =====")